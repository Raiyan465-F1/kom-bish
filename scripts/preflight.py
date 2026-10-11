"""Final pre-submission audit. Prints a PASS/FAIL table and never fixes anything by itself.

  python scripts/preflight.py [--no-network] [--demo-url URL ...]

Run it in .venv from the repo root. A FAIL names the file (and line) to look at.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WIRED = ["blb", "blast", "safety", "healthy", "not_sure", "consent", "voice_failed"]
DOIS = ["10.17632/kx9rx8p2mz.1", "10.17632/sjkptpgv79.2", "10.17632/znsxdctwtt.1"]
PINS = ["gradio==6.29.1", "onnxruntime==1.31.0", "pillow==12.3.0", "huggingface_hub==2.2.0", "google-genai==2.29.0"]
PLACEHOLDERS = ["TODO-C", "TODO-HUMAN", "VIDEO_URL_TODO", "DEMO_URL_TODO", "<gh-user>", "<hf-user>"]
SECRET_PATTERNS = [re.compile(r"hf_[A-Za-z0-9]{20,}"), re.compile(r"AIza[0-9A-Za-z_\-]{20,}"),
                   re.compile(r"AQ\.[0-9A-Za-z_\-]{20,}"),   # HF tokens and both Gemini key formats
                   re.compile(r"KGAT_[0-9A-Za-z]{20,}")]      # Kaggle API tokens
HAND_CHECK_HOSTS = ("youtube.com", "youtu.be", "drive.google.com", "docs.google.com")

results: list[tuple[str, str, str]] = []


def record(status: str, name: str, detail: str = "") -> None:
    results.append((status, name, detail))


def check(name: str, ok: bool, detail: str = "") -> bool:
    record("PASS" if ok else "FAIL", name, "" if ok else detail)
    return ok


def git(*args: str) -> str:
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
                          errors="replace").stdout.strip()


def tracked_files() -> list[str]:
    return [f for f in git("ls-files").splitlines() if f]


def read(rel: str) -> str:
    try:
        return (ROOT / rel).read_text("utf-8")
    except (UnicodeDecodeError, FileNotFoundError):
        return ""


# Code that names a placeholder token on purpose (a check for it, or a default the human overrides).
PLACEHOLDER_CODE = {"scripts/preflight.py", "scripts/publish_model.py", "deck/make_deck.py", "tests/test_rules.py"}


def exempt_from_placeholders(rel: str) -> bool:
    name = Path(rel).name
    return (rel.startswith("steps/") or rel in ("CLAUDE.md", ".env.example") or rel in PLACEHOLDER_CODE
            or "_template" in name or "_draft" in name)


def http_status(url: str, timeout: int = 90) -> int:
    req = urllib.request.Request(url, headers={"User-Agent": "kombish-preflight"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status
    except urllib.error.HTTPError as e:
        return e.code
    except Exception:
        return 0


def check_git() -> None:
    check("git: on main", git("rev-parse", "--abbrev-ref", "HEAD") == "main", "checkout main")
    dirty = git("status", "--short")
    check("git: clean tree", not dirty, "uncommitted: " + ", ".join(dirty.splitlines()[:5]))
    check("git: v1.0 (or later) tag exists", any(t.startswith("v1.") for t in git("tag").splitlines()),
          "tag v1.0 first (step 21)")
    subprocess.run(["git", "fetch", "--quiet", "origin"], cwd=ROOT)
    check("git: main equals origin/main", git("rev-parse", "main") == git("rev-parse", "origin/main"),
          "push or pull so local main matches origin/main")


def check_tests() -> None:
    for name, extra in (("pytest -q", []), ("pytest -q -m final", ["-m", "final"])):
        run = subprocess.run([sys.executable, "-m", "pytest", "-q", *extra], cwd=ROOT, capture_output=True,
                             text=True, encoding="utf-8", errors="replace",
                             env=os.environ | {"PYTHONIOENCODING": "utf-8"})
        tail = (run.stdout.strip().splitlines() or [""])[-1]
        failed = "; ".join(line for line in run.stdout.splitlines() if line.startswith("FAILED"))
        check(f"tests: {name}", run.returncode == 0, f"{tail} {failed}".strip())


def check_secrets(files: list[str]) -> None:
    check("secrets: .env not tracked", ".env" not in files, ".env is tracked: rotate every key in it, then untrack it")
    check("secrets: submission/private/ not tracked", not any(f.startswith("submission/private/") for f in files),
          "untrack submission/private/")
    hits = []
    for rel in files:
        for n, line in enumerate(read(rel).splitlines(), 1):
            if any(p.search(line) for p in SECRET_PATTERNS):
                hits.append(f"{rel}:{n}")
    check("secrets: no token-shaped strings in tracked files", not hits,
          "ROTATE these keys now, then remove them: " + ", ".join(hits[:10]))


def check_placeholders(files: list[str]) -> None:
    hits = []
    for rel in files:
        if exempt_from_placeholders(rel):
            continue
        for n, line in enumerate(read(rel).splitlines(), 1):
            for token in PLACEHOLDERS:
                if token in line:
                    hits.append(f"{rel}:{n} ({token})")
    more = f" ... and {len(hits) - 12} more" if len(hits) > 12 else ""
    check("placeholders: none left in tracked files", not hits, ", ".join(hits[:12]) + more)


def check_readme(demo_urls: list[str]) -> str:
    text = read("README.md")
    check("README: starts with YAML front matter", text.startswith("---\n"), "README.md:1")
    for needle in ('sdk_version: 6.29.1', 'python_version: "3.12"', "```mermaid", "Real vs mocked", "CC BY 4.0", *DOIS):
        check(f"README: contains {needle!r}", needle in text, "README.md")
    check("README: has a YouTube URL", bool(re.search(r"https://(www\.)?(youtube\.com|youtu\.be)/", text)), "README.md")
    for url in demo_urls:
        check(f"README: has demo URL {url}", url in text, "README.md")
    if not demo_urls:
        record("FAIL", "README: demo URL", "pass --demo-url <https URL> so it can be checked")
    check("README: under 500 KiB", len(text.encode("utf-8")) < 500 * 1024, "README.md")
    return text


def check_numbers(readme: str) -> None:
    path = ROOT / "training" / "eval_summary.json"
    if not path.exists():
        record("FAIL", "numbers: training/eval_summary.json exists", "run step 20")
        return
    summary = json.loads(path.read_text("utf-8"))
    missing = [k for k, v in summary.items() if v is None]
    check("numbers: eval_summary.json has every value", not missing, "null: " + ", ".join(missing))
    # A README number must be one of the summary's values, in one of the usual forms.
    forms: set[str] = set()
    for v in summary.values():
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            continue
        if isinstance(v, int):
            forms.add(str(v))
        else:
            forms.update({f"{v:.2f}", f"{v:.3f}", f"{v:.4f}", f"{100 * v:.0f}", f"{100 * v:.1f}"})
    section = re.search(r"## Evaluation\n(.*?)(?=\n## )", readme, flags=re.S)
    quoted = set(re.findall(r"(?<![\w./-])\d+(?:\.\d+)?(?![\w/-])", section.group(1))) if section else set()
    unknown = sorted(n for n in quoted if n not in forms and n not in {"95", "16123"})
    check("numbers: every number in README 'Evaluation' is in eval_summary.json", bool(section) and not unknown,
          "README.md 'Evaluation': " + (", ".join(unknown) if section else "section not found"))


def check_requirements() -> None:
    full = [line.strip() for line in read("requirements.txt").splitlines() if line.strip()]
    render = [line.strip() for line in read("requirements-render.txt").splitlines() if line.strip()]
    for name, lines in (("requirements.txt", full), ("requirements-render.txt", render)):
        absent = [p for p in PINS if p not in lines]
        check(f"requirements: pins in {name}", not absent, f"{name}: missing {absent}")
        check(f"requirements: no torch in {name}", not any(line.lower().startswith("torch") for line in lines), name)
    check("requirements: spaces only in requirements.txt", "spaces" in full and "spaces" not in render,
          "requirements.txt / requirements-render.txt")
    check("requirements: files differ only by spaces", [line for line in full if line != "spaces"] == render,
          "requirements.txt vs requirements-render.txt")


def check_files() -> None:
    from PIL import Image

    missing = [c for c in WIRED if not ((ROOT / "audio" / f"{c}.mp3").exists()
                                        and (ROOT / "audio" / f"{c}.mp3").stat().st_size > 0)]
    check("files: 7 wired audio clips", not missing, "missing audio/: " + ", ".join(f"{c}.mp3" for c in missing))
    examples = sorted((ROOT / "examples").glob("*.jpg"))
    check("files: 4-6 example JPEGs", 4 <= len(examples) <= 6, f"examples/ has {len(examples)} JPEGs")
    with_exif = []
    for path in examples:
        with Image.open(path) as im:
            if len(im.getexif()):
                with_exif.append(path.name)
    check("files: examples have no EXIF", not with_exif, "run scripts/strip_exif.py on: " + ", ".join(with_exif))
    deck = ROOT / "deck" / "Kom_Bish_deck.pptx"
    if not deck.exists():
        record("FAIL", "files: deck exists", "deck/Kom_Bish_deck.pptx (step 22)")
    else:
        try:
            from pptx import Presentation
            n = len(Presentation(str(deck)).slides)
            check("files: deck has 8 slides or fewer", n <= 8, f"deck/Kom_Bish_deck.pptx has {n} slides")
        except ImportError:
            record("HAND", "files: deck slide count", "python-pptx is not in .venv; count the slides by hand")


def check_links(demo_urls: list[str]) -> None:
    origin = git("remote", "get-url", "origin").removesuffix(".git")
    check(f"links: GitHub repo is public ({origin})", http_status(origin, 30) == 200,
          "repo returns an error without login: make it public (Settings > General > Danger Zone)")
    for url in demo_urls:
        check(f"links: demo responds ({url})", http_status(url, 90) == 200, f"{url} did not return 200 within 90 s")
    try:
        from dotenv import dotenv_values
        repo = (os.getenv("MODEL_REPO") or dotenv_values(ROOT / ".env").get("MODEL_REPO") or "").strip()
    except ImportError:
        repo = os.getenv("MODEL_REPO", "").strip()
    if not repo or "<" in repo:
        record("FAIL", "links: MODEL_REPO downloads with no token", "MODEL_REPO is not set in .env (step 11)")
    else:
        code = ("import os, sys; os.environ.pop('HF_TOKEN', None); from huggingface_hub import hf_hub_download; "
                "hf_hub_download(repo_id=sys.argv[1], filename='labels.json', token=False)")
        with tempfile.TemporaryDirectory() as cache:
            env = {k: v for k, v in os.environ.items() if k != "HF_TOKEN"} | {"HF_HUB_CACHE": cache}
            ok = subprocess.run([sys.executable, "-c", code, repo], env=env, capture_output=True).returncode == 0
        check(f"links: {repo} labels.json downloads with no token", ok, "model repo is private or missing")
    links = ROOT / "submission" / "private" / "links.txt"
    if not links.exists():
        record("FAIL", "links: submission/private/links.txt exists", "finish step 25")
        return
    for url in sorted(set(re.findall(r"https://[^\s<>()]+", links.read_text("utf-8")))):
        if any(host in url for host in HAND_CHECK_HOSTS):
            record("HAND", f"links: {url}", "check by hand in a private window")
        else:
            status = http_status(url, 90)
            check(f"links: {url}", 0 < status < 400, f"submission/private/links.txt: {url} returned {status or 'no response'}")


MANUAL = """
Manual checks (a human, not B, from a phone on mobile data and from a private laptop window):
  [ ] GitHub repo opens without login; README renders with the diagram; README links work.
  [ ] Demo wakes and completes one consultation (consent, photo, answer, audio plays).
  [ ] YouTube video plays (Unlisted link) with subtitles.
  [ ] Drive folder opens; the deck downloads; the video plays or shows processing.
  [ ] Model repo page opens.
  [ ] Write "Link audit passed - <name> - <time>" in docs/gates.md and commit (last commit before the email).
"""


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--demo-url", action="append", default=[], help="hosted demo URL (repeat for the mirror)")
    p.add_argument("--no-network", action="store_true", help="skip the link checks")
    p.add_argument("--skip-tests", action="store_true", help="skip the two pytest runs")
    args = p.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # Windows consoles default to a legacy code page

    files = tracked_files()
    check_git()
    if not args.skip_tests:
        check_tests()
    check_secrets(files)
    check_placeholders(files)
    readme = check_readme(args.demo_url)
    check_numbers(readme)
    check_requirements()
    check_files()
    if not args.no_network:
        check_links(args.demo_url)

    width = max(len(name) for _, name, _ in results)
    for status, name, detail in results:
        print(f"{status:<5} {name:<{width}}  {detail}")
    fails = sum(status == "FAIL" for status, _, _ in results)
    print(f"\n{len(results) - fails} of {len(results)} automated checks did not fail; {fails} FAIL.")
    print(MANUAL)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
